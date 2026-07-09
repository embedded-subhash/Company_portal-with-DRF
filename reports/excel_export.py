"""
Module 3 & 4: Excel Import & Export
Handles bulk Excel employee import (with per-row validation) and export of
Employee / Department / Salary / Attendance data to .xlsx.
"""
import io
from datetime import datetime, date

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from employees.models import Attendance, Department, Employee

REQUIRED_COLUMNS = [
    "Employee ID", "First Name", "Last Name", "Email", "Phone",
    "Department", "Designation", "Salary", "Joining Date",
]

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)


# ---------------------------------------------------------------------------
# Module 3: Excel Import
# ---------------------------------------------------------------------------
def import_employees_from_excel(file_obj):
    """
    Reads an uploaded .xlsx of employees, validates each row, skips invalid
    records (collecting their errors) and saves valid employees.
    Returns: {"total_rows": int, "created": int, "failed": int, "errors": [...]}
    """
    try:
        wb = load_workbook(file_obj, data_only=True)
    except Exception as exc:  # noqa: BLE001
        return {"total_rows": 0, "created": 0, "failed": 0, "errors": [f"Could not read Excel file: {exc}"]}

    ws = wb.active
    header_row = [cell.value for cell in ws[1]]
    missing = [c for c in REQUIRED_COLUMNS if c not in header_row]
    if missing:
        return {
            "total_rows": 0, "created": 0, "failed": 0,
            "errors": [f"Missing required column(s): {', '.join(missing)}"],
        }
    col_index = {name: header_row.index(name) for name in REQUIRED_COLUMNS}

    total_rows = 0
    created = 0
    errors = []
    seen_ids_in_file = set()
    seen_emails_in_file = set()

    for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row is None or all(v is None for v in row):
            continue
        total_rows += 1

        def cell(name):
            return row[col_index[name]]

        emp_id = str(cell("Employee ID") or "").strip()
        email = str(cell("Email") or "").strip().lower()

        if not emp_id or not email:
            errors.append(f"Row {row_num}: Employee ID and Email are required.")
            continue
        if emp_id in seen_ids_in_file or email in seen_emails_in_file:
            errors.append(f"Row {row_num}: Duplicate Employee ID/Email within file ({emp_id}).")
            continue
        if Employee.objects.filter(employee_id=emp_id).exists():
            errors.append(f"Row {row_num}: Employee ID '{emp_id}' already exists.")
            continue
        if Employee.objects.filter(email=email).exists():
            errors.append(f"Row {row_num}: Email '{email}' already exists.")
            continue

        try:
            salary = float(cell("Salary") or 0)
            raw_date = cell("Joining Date")
            if isinstance(raw_date, (datetime, date)):
                joining_date = raw_date if isinstance(raw_date, date) and not isinstance(raw_date, datetime) else raw_date.date()
            else:
                joining_date = datetime.strptime(str(raw_date).strip(), "%Y-%m-%d").date()

            dept_name = str(cell("Department") or "").strip()
            department, _ = Department.objects.get_or_create(
                name=dept_name, defaults={"code": dept_name[:20].upper() or "GEN"}
            )
            Employee.objects.create(
                employee_id=emp_id,
                first_name=str(cell("First Name") or "").strip(),
                last_name=str(cell("Last Name") or "").strip(),
                email=email,
                phone=str(cell("Phone") or "").strip(),
                department=department,
                designation=str(cell("Designation") or "").strip(),
                salary=salary,
                joining_date=joining_date,
            )
            seen_ids_in_file.add(emp_id)
            seen_emails_in_file.add(email)
            created += 1
        except Exception as exc:  # noqa: BLE001
            errors.append(f"Row {row_num}: {exc}")

    return {
        "total_rows": total_rows,
        "created": created,
        "failed": total_rows - created,
        "errors": errors,
    }


# ---------------------------------------------------------------------------
# Module 4: Excel Export
# ---------------------------------------------------------------------------
def _style_header(ws, headers):
    for idx, header in enumerate(headers, start=1):
        c = ws.cell(row=1, column=idx, value=header)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        ws.column_dimensions[get_column_letter(idx)].width = max(14, len(header) + 4)


def _workbook_to_bytes(wb):
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def export_employee_list_excel(queryset):
    wb = Workbook()
    ws = wb.active
    ws.title = "Employees"
    headers = REQUIRED_COLUMNS + ["Status"]
    _style_header(ws, headers)
    for emp in queryset:
        ws.append([
            emp.employee_id, emp.first_name, emp.last_name, emp.email, emp.phone,
            emp.department.name if emp.department else "", emp.designation,
            float(emp.salary), emp.joining_date.isoformat(), emp.status,
        ])
    return _workbook_to_bytes(wb)


def export_department_list_excel(queryset):
    wb = Workbook()
    ws = wb.active
    ws.title = "Departments"
    _style_header(ws, ["Department", "Code", "Employee Count"])
    for dept in queryset:
        ws.append([dept.name, dept.code, dept.employees.count()])
    return _workbook_to_bytes(wb)


def export_salary_report_excel(queryset):
    wb = Workbook()
    ws = wb.active
    ws.title = "Salary Report"
    _style_header(ws, ["Employee ID", "Name", "Department", "Designation", "Salary"])
    for emp in queryset:
        ws.append([
            emp.employee_id, emp.full_name,
            emp.department.name if emp.department else "", emp.designation, float(emp.salary),
        ])
    return _workbook_to_bytes(wb)


def export_attendance_report_excel(queryset):
    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance Report"
    _style_header(ws, ["Employee ID", "Name", "Date", "Status"])
    for record in queryset.select_related("employee"):
        ws.append([
            record.employee.employee_id, record.employee.full_name,
            record.date.isoformat(), record.status,
        ])
    return _workbook_to_bytes(wb)
